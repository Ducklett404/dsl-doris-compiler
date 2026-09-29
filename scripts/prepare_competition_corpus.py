#!/usr/bin/env python3
"""Prepare the organizer's stock-screening DSL workbook without changing it.

Requires openpyxl. Example:
  python scripts/prepare_competition_corpus.py --input ".../选股提槽实例.xlsx"
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


FUNCTION_ARITY = {
    "indicator": 7,
    "rank": 8,
    "belong_industry": 5,
    "queryData": 5,
    "cross_above": 7,
    "cross_down": 7,
    "close_together": 7,
    "reset": 7,
    "golden_cross": 7,
    "record_high": 8,
    "record_low": 8,
    "position": 8,
}
CALL_PATTERN = re.compile(r"\b(" + "|".join(map(re.escape, FUNCTION_ARITY)) + r")\s*\(")
BADCASE_PATTERN = re.compile(r"bad\s*case|badcase|错误|待修", re.IGNORECASE)
WHITESPACE_PATTERN = re.compile(r"\s+")


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFC", str(value)).replace("\r\n", "\n").replace("\r", "\n").strip()


def normalize_query(value: object) -> str:
    return WHITESPACE_PATTERN.sub(" ", normalize_text(value))


def dsl_identity(value: str) -> str:
    """Ignore formatting outside literals, but never alter quoted values."""
    source = normalize_text(value)
    out: list[str] = []
    quoted = False
    pending_space = False
    index = 0
    while index < len(source):
        char = source[index]
        if char == "'":
            if pending_space and out:
                out.append(" ")
            pending_space = False
            out.append(char)
            if quoted and index + 1 < len(source) and source[index + 1] == "'":
                out.append("'")
                index += 2
                continue
            quoted = not quoted
        elif not quoted and char.isspace():
            pending_space = True
        else:
            if pending_space and out:
                out.append(" ")
            pending_space = False
            out.append(char)
        index += 1
    return "".join(out)


def parse_call_arguments(source: str, open_paren: int) -> tuple[list[str], int]:
    arguments: list[str] = []
    depth = 1
    quoted = False
    start = open_paren + 1
    index = start
    while index < len(source):
        char = source[index]
        if char == "'":
            if quoted and index + 1 < len(source) and source[index + 1] == "'":
                index += 2
                continue
            quoted = not quoted
        elif not quoted:
            if char == "(":
                depth += 1
            elif char == ")":
                depth -= 1
                if depth == 0:
                    final = source[start:index].strip()
                    if arguments or final:
                        arguments.append(final)
                    return arguments, index + 1
            elif char == "," and depth == 1:
                arguments.append(source[start:index].strip())
                start = index + 1
        index += 1
    raise ValueError("unclosed function call or string literal")


def scan_calls(source: str) -> tuple[list[tuple[str, list[str]]], list[str]]:
    calls: list[tuple[str, list[str]]] = []
    warnings: list[str] = []
    for match in CALL_PATTERN.finditer(source):
        function = match.group(1)
        try:
            args, _ = parse_call_arguments(source, match.end() - 1)
        except ValueError as error:
            warnings.append(f"{function}: {error}")
            continue
        calls.append((function, args))
        if len(args) != FUNCTION_ARITY[function]:
            warnings.append(f"{function}: expected {FUNCTION_ARITY[function]} args, got {len(args)}")
    return calls, warnings


def quoted_literal(argument: str) -> str | None:
    argument = argument.strip()
    if len(argument) < 2 or argument[0] != "'" or argument[-1] != "'":
        return None
    inner = argument[1:-1]
    if "'" in inner.replace("''", ""):
        return None
    return inner.replace("''", "'")


def split_for(query: str) -> str:
    # Same normalized question always lands in the same subset.
    bucket = int(hashlib.sha256(query.encode("utf-8")).hexdigest()[:8], 16) % 10
    return "train" if bucket < 8 else "dev" if bucket == 8 else "test"


def classify_rows(source_rows: list[tuple[int, object, object, object, object]]) -> list[dict]:
    records: list[dict] = []
    for row_number, batch, query, dsl, note in source_rows:
        records.append({
            "source_row": row_number,
            "batch": normalize_text(batch),
            "query": normalize_query(query),
            "dsl": normalize_text(dsl),
            "note": normalize_text(note),
        })

    variants: dict[str, set[str]] = defaultdict(set)
    for record in records:
        if record["query"] and record["dsl"] and not BADCASE_PATTERN.search(record["note"]):
            variants[record["query"]].add(dsl_identity(record["dsl"]))
    conflicts = {query for query, forms in variants.items() if len(forms) > 1}
    seen: set[tuple[str, str]] = set()

    for record in records:
        query, dsl, note = record["query"], record["dsl"], record["note"]
        calls, warnings = scan_calls(dsl) if dsl else ([], [])
        record["function_calls"] = [name for name, _ in calls]
        record["scan_warnings"] = warnings
        if not query:
            status, reason = "excluded", "missing_query"
        elif not dsl:
            status, reason = "excluded", "missing_dsl"
        elif BADCASE_PATTERN.search(note):
            status, reason = "review", "badcase_note"
        elif query in conflicts:
            status, reason = "review", "conflicting_dsl_for_query"
        elif warnings:
            status, reason = "review", "function_scan_warning"
        elif (query, dsl_identity(dsl)) in seen:
            status, reason = "excluded", "exact_duplicate"
        else:
            status, reason = "included", "unvalidated_label"
            seen.add((query, dsl_identity(dsl)))
        record["status"] = status
        record["reason"] = reason
        record["split"] = split_for(query) if status == "included" else ""
    return records


def extract_indicators(records: list[dict]) -> list[dict]:
    inventory: dict[str, dict] = {}
    for record in records:
        if record["status"] != "included":
            continue
        calls, _ = scan_calls(record["dsl"])
        for function, args in calls:
            indices = (2, 3) if function in {"cross_above", "cross_down", "close_together", "reset", "golden_cross"} else (2,)
            for index in indices:
                if index >= len(args):
                    continue
                name = quoted_literal(args[index])
                if not name:
                    continue
                item = inventory.setdefault(name, {"indicator": name, "mentions": 0, "functions": Counter(), "source_rows": []})
                item["mentions"] += 1
                item["functions"][function] += 1
                if len(item["source_rows"]) < 3 and record["source_row"] not in item["source_rows"]:
                    item["source_rows"].append(record["source_row"])
    return sorted(inventory.values(), key=lambda item: (-item["mentions"], item["indicator"]))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_outputs(input_path: Path, output_dir: Path) -> dict:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    if "训练语料" not in workbook.sheetnames:
        raise ValueError("expected sheet 训练语料")
    sheet = workbook["训练语料"]
    header = tuple(cell.value for cell in sheet[1][:4])
    if header != ("添加时间", "query", "slot_dsl", "备注"):
        raise ValueError(f"unexpected header: {header!r}")
    rows = [(index, *values[:4]) for index, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2)]
    records = classify_rows(rows)
    inventory = extract_indicators(records)
    output_dir.mkdir(parents=True, exist_ok=True)

    with (output_dir / "corpus.jsonl").open("w", encoding="utf-8", newline="\n") as target:
        for record in records:
            if record["status"] == "included":
                target.write(json.dumps({key: record[key] for key in ("source_row", "batch", "query", "dsl", "note", "split")}, ensure_ascii=False) + "\n")

    with (output_dir / "row_audit.csv").open("w", encoding="utf-8-sig", newline="") as target:
        writer = csv.writer(target)
        writer.writerow(("source_row", "status", "reason", "split", "function_calls", "scan_warnings"))
        for record in records:
            writer.writerow((record["source_row"], record["status"], record["reason"], record["split"],
                             ";".join(record["function_calls"]), ";".join(record["scan_warnings"])))

    with (output_dir / "review_queue.csv").open("w", encoding="utf-8-sig", newline="") as target:
        writer = csv.writer(target)
        writer.writerow(("source_row", "reason", "query", "dsl", "note", "scan_warnings"))
        for record in records:
            if record["status"] != "included":
                writer.writerow((record["source_row"], record["reason"], record["query"], record["dsl"],
                                 record["note"], ";".join(record["scan_warnings"])))

    with (output_dir / "indicator_inventory.csv").open("w", encoding="utf-8-sig", newline="") as target:
        writer = csv.writer(target)
        writer.writerow(("indicator", "mentions", "functions", "example_source_rows", "mapping_status"))
        for item in inventory:
            writer.writerow((item["indicator"], item["mentions"],
                             ";".join(f"{name}:{count}" for name, count in sorted(item["functions"].items())),
                             ";".join(map(str, item["source_rows"])), "UNMAPPED"))

    status_counts = Counter(record["status"] for record in records)
    reason_counts = Counter(record["reason"] for record in records)
    split_counts = Counter(record["split"] for record in records if record["status"] == "included")
    function_counts = Counter(name for record in records if record["status"] == "included" for name in record["function_calls"])
    manifest = {
        "schema_version": 1,
        "source_filename": input_path.name,
        "source_sha256": sha256_file(input_path),
        "sheet": "训练语料",
        "header": header,
        "source_rows": len(records),
        "status_counts": dict(sorted(status_counts.items())),
        "reason_counts": dict(sorted(reason_counts.items())),
        "split_counts": dict(sorted(split_counts.items())),
        "function_counts_included": dict(sorted(function_counts.items())),
        "unique_indicators_included": len(inventory),
        "normalization": "NFC; CRLF/CR to LF; trim outer whitespace; collapse query whitespace only; DSL interior preserved",
        "validation_scope": "format and duplicate screening only; no DSL semantic or SQL correctness validation",
    }
    (output_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    workbook.close()
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="organizer's 选股提槽实例.xlsx")
    parser.add_argument("--output-dir", type=Path, default=Path("build/competition-corpus"))
    args = parser.parse_args()
    if not args.input.is_file():
        parser.error(f"input workbook not found: {args.input}")
    print(json.dumps(write_outputs(args.input, args.output_dir), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
